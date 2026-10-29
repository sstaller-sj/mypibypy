"""The ranking benchmark: reading ordered designs from workbooks, and ranking them.

Workbooks here are synthetic, built in both layouts the lab's order workbooks
use, so the reader is tested without the lab's own files.
"""
import random
import re

import pytest
from openpyxl import Workbook

from mypibypy.benchmark import markdown, rank_design, read_workbook
from mypibypy.design import FOOTPRINT, LEFT_LEN, GAP_LEN, scan
from mypibypy.sequence import GenomicInterval

DNA_RUN = re.compile(r"[ACGT]{12,}")


def _seq(n, seed=5):
    rng = random.Random(seed)
    return "".join(rng.choice("ACGT") for _ in range(n))


def _pads(seq, offset):
    return (seq[offset:offset + LEFT_LEN],
            seq[offset + LEFT_LEN + GAP_LEN:offset + FOOTPRINT])


def _single_sheet(wb, title, region_seq, start, offset, corrupt=False):
    ws = wb.create_sheet(title)
    left, right = _pads(region_seq, offset)
    if corrupt:
        left = "A" * 25
    end = start + len(region_seq) - 1
    fp = start + offset
    for row in (["Channel", "Cy5 (system essrb_1)"],
                ["Target region", f"mm10 chr7:{start:,}-{end:,}"],
                ["Chosen footprint", f"mm10 chr7:{fp:,}-{fp + FOOTPRINT - 1:,}"],
                ["left25 (plus)", left], ["gap (skipped)", "G"],
                ["right25 (plus)", right],
                ["Region sequence (plus, 5'->3')", region_seq]):
        ws.append(row)


def _tiled_sheet(wb, title, region_seq, offsets):
    ws = wb.create_sheet(title)
    ws.append([f"{title} - test"])
    ws.append(["2 tiled pairs into Cy5 (system essrb_1); target region ENST0000001 transcript"])
    ws.append([])
    ws.append(["#", "Footprint", "left25 (plus)", "gap", "right25 (plus)", "GC L/R"])
    for i, off in enumerate(offsets, 1):
        left, right = _pads(region_seq, off)
        ws.append([str(i), f"ENST0000001 transcript:{off + 1}-{off + FOOTPRINT}",
                   left, "T", right, "50%/50%"])
    ws.append([])
    ws.append(["Region sequence (plus, 5'->3')", region_seq])


@pytest.fixture
def genomic():
    seq = _seq(300)
    best = scan(seq, GenomicInterval("mm10", "chr7", 1000, 1299), top=1)[0]
    return seq, best.offset


@pytest.fixture
def workbook(tmp_path, genomic):
    seq, best = genomic
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("ORDER_ssDNA").append(["NAME", "SEQUENCE"])
    _single_sheet(wb, "Locus_A", seq, 1000, best)
    _tiled_sheet(wb, "Tx_sense", _seq(400, seed=9), [10, 200])
    fusion = wb.create_sheet("FUSION")                # no footprint to rank
    fusion.append(["SENSE TARGET 5->3", _seq(51)])
    path = tmp_path / "order.xlsx"
    wb.save(path)
    return path


def test_both_layouts_are_read_and_fusions_skipped(workbook):
    designs = {d.sheet: d for d in read_workbook(workbook)}
    assert set(designs) == {"Locus_A", "Tx_sense"}
    locus, tx = designs["Locus_A"], designs["Tx_sense"]
    assert locus.kind == "genomic" and locus.interval.start == 1000
    assert tx.kind == "rna" and tx.accession == "ENST0000001"
    assert [p.offset for p in tx.pairs] == [10, 200]
    assert not locus.problems and not tx.problems


def test_offsets_are_zero_based_within_the_region(workbook, genomic):
    _, best = genomic
    locus = next(d for d in read_workbook(workbook) if d.sheet == "Locus_A")
    assert locus.pairs[0].offset == best


def test_a_pad_not_at_its_footprint_is_reported(tmp_path, genomic):
    seq, best = genomic
    wb = Workbook()
    wb.remove(wb.active)
    _single_sheet(wb, "Bad", seq, 1000, best, corrupt=True)
    wb.save(tmp_path / "bad.xlsx")
    (design,) = read_workbook(tmp_path / "bad.xlsx")
    assert design.problems == ["Bad: left pad is not at its footprint"]


def test_the_pipelines_own_pick_ranks_first(workbook):
    locus = next(d for d in read_workbook(workbook) if d.sheet == "Locus_A")
    (row,) = rank_design(locus, "v1")
    assert row.rank == 1 and row.picked and row.cost == row.best_cost
    assert row.windows == len(locus.seq) - FOOTPRINT + 1


def test_a_masked_footprint_is_not_a_candidate(workbook):
    locus = next(d for d in read_workbook(workbook) if d.sheet == "Locus_A")
    fp = locus.interval.start + locus.pairs[0].offset
    rmsk = [{"genoStart": fp - 1, "genoEnd": fp + 10}]
    (row,) = rank_design(locus, "v1", rmsk)
    assert row.rank is None and not row.picked and "not a candidate" in row.note


def test_a_tiled_design_ranks_every_pair(workbook):
    tx = next(d for d in read_workbook(workbook) if d.sheet == "Tx_sense")
    rows = rank_design(tx, "v1")
    assert [r.pair for r in rows] == ["Tx_sense #1", "Tx_sense #2"]
    assert all(r.rank is not None for r in rows)


def test_the_report_carries_no_sequence_and_no_names_by_default(workbook):
    rows = [r for d in read_workbook(workbook) for r in rank_design(d, "v1")]
    table = markdown(rows, {"Locus_A": "worked"})
    assert not DNA_RUN.search(table)
    assert "Locus_A" not in table and "Tx_sense" not in table
    assert "worked" in table                       # outcomes still apply by label
    assert "Locus_A" in markdown(rows, names=True)
