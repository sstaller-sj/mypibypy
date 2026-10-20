"""Probe self-folding via RNAfold with DNA parameters.

Synthetic oligos only. The structural cases are built so the answer is not in
doubt -- a strong designed hairpin and a poly-T that cannot pair -- because the
point is that the plumbing reports the right oligo, region and status, not that
ViennaRNA's energies are right.
"""
from __future__ import annotations

import random
import shutil
from dataclasses import dataclass

import pytest
from openpyxl import load_workbook

from mypibypy import fold
from mypibypy.design import build_pair, scan
from mypibypy.fold import (
    EVALUATED,
    NOT_EVALUATED,
    OligoFold,
    PairFolding,
    describe_folding,
    fold_oligos,
    pair_folding,
    parse_dot_plot,
    parse_mfe_frequencies,
    parse_rnafold,
)
from mypibypy.loci import Locus
from mypibypy.pipeline import DesignResult
from mypibypy.sequence import GenomicInterval, revcomp
from mypibypy.specificity import not_screened
from mypibypy.systems import systems
from mypibypy.workbook import write_workbook

HAVE_RNAFOLD = bool(shutil.which("RNAfold"))
needs_rnafold = pytest.mark.skipif(not HAVE_RNAFOLD, reason="no RNAfold")

STEM = "GCGCAGCCGC"
HAIRPIN = "AAAA" + STEM + "TTTT" + revcomp(STEM) + "AAAA"


class TestParsing:
    def test_parses_structure_and_energy(self):
        out = ">o0\nGGGG\n((((....)))) ( -3.20)\n>o1\nAAAA\n.... (  0.00)\n"
        assert parse_rnafold(out) == [("((((....))))", -3.2), ("....", 0.0)]

    def test_only_the_mfe_line_is_taken_from_partition_function_output(self):
        out = (">o0\nGGGG\n((((....)))) ( -3.20)\n((((,...)))) [ -3.50]\n"
               "((((....)))) { -3.20 d=0.50}\n frequency of mfe structure in "
               "ensemble 0.62; ensemble diversity 0.9\n")
        assert parse_rnafold(out) == [("((((....))))", -3.2)]
        assert parse_mfe_frequencies(out) == [0.62]

    def test_dot_plot_sums_partner_probabilities(self):
        dp = "%comment\n1 4 0.5 ubox\n1 3 0.5 ubox\n2 3 0.9 lbox\n"
        probs = parse_dot_plot(dp, 4)
        assert probs == pytest.approx((0.5, 0.0, 0.25, 0.25))

    def test_ensemble_probabilities_take_precedence_over_the_mfe(self):
        o = OligoFold("x", "N" * 4, -1.0, "(())", pad=(0, 4), arm=(0, 0),
                      paired_prob=(0.1, 0.1, 0.1, 0.1))
        assert o.pad_paired == pytest.approx(0.1)

    def test_paired_fractions_by_region(self):
        o = OligoFold("x", "N" * 10, -2.0, "((....))..", pad=(5, 10), arm=(0, 5))
        assert o.arm_paired == pytest.approx(2 / 5)
        assert o.pad_paired == pytest.approx(2 / 5)

    def test_words_are_bands_not_cutoffs(self):
        assert describe_folding(0.0) == "open"
        assert describe_folding(0.45) == "open", "random-sequence background"
        assert describe_folding(0.55) == "partly folded"
        assert describe_folding(0.8) == "folded"


@dataclass
class _Sys:
    arm_left: str
    arm_right: str


@dataclass
class _Pair:
    system: _Sys
    p1l_name: str
    p1r_name: str
    p1l: str
    p1r: str


class TestNotEvaluated:
    def test_missing_tool_is_not_evaluated_with_a_reason(self, monkeypatch):
        monkeypatch.setenv("MYPIBYPY_RNAFOLD", "/nonexistent/RNAfold")
        res = pair_folding(_Pair(_Sys("A" * 22, "T" * 22), "L", "R", "A" * 47, "T" * 47))
        assert res.status == NOT_EVALUATED and not res.evaluated
        assert "RNAfold" in res.reason
        assert res.summary().startswith("NOT EVALUATED")
        assert res.worst_paired is None and res.min_dg is None

    def test_a_hand_built_result_reports_not_evaluated(self):
        sys_a = list(systems().values())[0]
        rng = random.Random(7)
        seq = "".join(rng.choice("ACGT") for _ in range(300))
        iv = GenomicInterval("mm10", "chr1", 1, 300)
        cands = scan(seq, iv, top=1)
        res = DesignResult(Locus("p", "G", "", iv, sys_a.name), sys_a, seq, cands,
                           build_pair("p", sys_a, cands[0], "P1L", "P1R"), [],
                           not_screened("x"))
        assert not res.folding.evaluated
        assert not res.folding_at(3).evaluated


@needs_rnafold
class TestRnafold:
    def test_dna_parameters_are_used(self):
        _, tool = fold_oligos([HAIRPIN])
        assert "DNA" in tool

    def test_a_designed_hairpin_folds_and_poly_t_does_not(self):
        hp, pt = fold_oligos([HAIRPIN, "T" * 30])[0]
        assert hp.mfe_dg < -3 and hp.structure.count("(") >= 8
        assert hp.mfe_frequency > 0.5
        stem_bases = [i for i, c in enumerate(hp.structure) if c != "."]
        assert min(hp.paired_prob[i] for i in stem_bases) > 0.5
        assert pt.mfe_dg == 0.0 and set(pt.structure) == {"."}
        assert max(pt.paired_prob) < 0.01

    def test_a_weak_mfe_is_not_called_folded(self):
        """The case that motivated the ensemble metric: an MFE nobody populates."""
        weak = "GCATTGCTGACAGGATGCAGAAGGATGCAGCACGATAGGGAATCGGA"
        (res,), _ = fold_oligos([weak])
        mfe_paired = sum(c != "." for c in res.structure) / len(weak)
        ensemble_paired = sum(res.paired_prob) / len(weak)
        assert res.mfe_frequency < 0.2
        assert ensemble_paired < mfe_paired

    def test_the_folded_region_is_attributed_to_pad_or_arm(self):
        # Poly-A arm and C-rich loops/flanks: nothing for the arm to pair with,
        # so any structure is the pad's own hairpin.
        arm = "A" * 22
        pad = "CCCC" + STEM + "CCCC" + revcomp(STEM) + "CCCC"
        pair = _Pair(_Sys(arm, arm), "P1L_x", "P1R_x",
                     p1l=arm + pad, p1r="A" * 25 + arm)
        res = pair_folding(pair)
        assert res.status == EVALUATED
        assert res.p1l.pad_paired > 0.3 and res.p1l.arm_paired < 0.05
        assert res.p1r.worst_paired < 0.01
        assert res.limiting is res.p1l
        assert res.limiting.label != "open"  # 20 of 32 pad bases are stem
        assert "UNCALIBRATED" in res.summary()
        assert res.to_dict()["p1l"]["pad_paired"] == res.p1l.pad_paired

    def test_a_real_pair_from_the_configured_systems(self, tmp_path):
        sys_a = list(systems().values())[0]
        rng = random.Random(11)
        seq = "".join(rng.choice("ACGT") for _ in range(400))
        iv = GenomicInterval("mm10", "chr1", 1, 400)
        cands = scan(seq, iv, top=1)
        pair = build_pair("p", sys_a, cands[0], "P1L_G_p", "P1R_G_p")
        res = pair_folding(pair)
        assert res.evaluated
        assert res.p1l.pad == (len(sys_a.arm_left), len(pair.p1l))
        assert res.p1r.arm == (len(pair.p1r) - len(sys_a.arm_right), len(pair.p1r))

        result = DesignResult(Locus("p", "G", "", iv, sys_a.name), sys_a, seq, cands,
                              pair, [], not_screened("x"), pair_folding=[res])
        out = tmp_path / "o.xlsx"
        write_workbook(str(out), [result], "test")
        ws = load_workbook(out)["G_p"]
        assert ws["A15"].value == "Probe self-folding"
        assert "P1L_G_p" in ws["B15"].value and "ensemble" in ws["B15"].value
        notes = load_workbook(out)["DESIGN_NOTES"]
        assert any(c.value == "Probe folding" for c in notes["A"])
