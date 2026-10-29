"""Tiling: several non-overlapping pairs per target.

One pair at ~100x amplification already works and every validated design in the
lab is single-pair, so tiling is headroom for marginal signal rather than a fix.
That shapes these tests: the single-pair path must come out exactly as it did
before tiling existed, and `TestSinglePairIsUnchanged` is the guard for it.

The other thing worth holding down is that *every* selected pair is screened.
Screening only the best one and ordering ten would be worse than not tiling,
because the unchecked nine look as validated as the first.
"""
from __future__ import annotations

import random
from itertools import combinations

import pytest
from openpyxl import load_workbook

from mypibypy.design import (
    DEFAULT_MIN_SPACING_NT,
    FOOTPRINT,
    build_pair,
    build_tiled_pairs,
    repack,
    scan,
    select_tiled,
    tile,
    tiled_names,
)
from mypibypy.loci import Locus
from mypibypy.pipeline import DesignResult, FusionIsSinglePair, design_fusion, design_locus
from mypibypy.sequence import GenomicInterval
from mypibypy.specificity import not_screened
from mypibypy.systems import systems
from mypibypy.workbook import write_workbook


@pytest.fixture(scope="module")
def sys_a():
    reg = systems()
    assert reg, "no amplifier systems configured"
    return list(reg.values())[0]


@pytest.fixture(scope="module")
def region():
    rng = random.Random(1234)
    seq = "".join(rng.choice("ACGT") for _ in range(600))
    return seq, GenomicInterval("mm10", "chr1", 2_000_000, 2_000_599)


@pytest.fixture(scope="module")
def ranked(region):
    seq, iv = region
    return scan(seq, iv, top=len(seq))


def make_result(sys_a, region, n_pairs, screens=None):
    seq, iv = region
    locus = Locus("prom", "Gene1", "synthetic", iv, sys_a.name)
    cands = scan(seq, iv, top=len(seq))
    selected = select_tiled(cands, n_pairs)
    pairs = build_tiled_pairs(locus.name, locus.gene, sys_a, selected)
    screens = screens or [not_screened("test") for _ in pairs]
    return DesignResult(locus, sys_a, seq, cands[:5], pairs[0], [], screens[0],
                        pairs=pairs, pair_screens=screens)


class TestSelection:
    def test_single_pair_returns_exactly_the_best_candidate(self, ranked):
        assert select_tiled(ranked, 1) == [ranked[0]]

    def test_selected_footprints_never_overlap(self, ranked):
        for n in (2, 4, 8):
            offsets = sorted(c.offset for c in select_tiled(ranked, n))
            for a, b in zip(offsets, offsets[1:]):
                assert b - a >= FOOTPRINT, f"footprints at {a} and {b} overlap"

    def test_minimum_spacing_is_respected(self, ranked):
        for spacing in (0, 3, 50):
            offsets = sorted(c.offset for c in select_tiled(ranked, 5, spacing))
            for a, b in zip(offsets, offsets[1:]):
                assert b - a >= FOOTPRINT + spacing

    def test_wider_spacing_fits_fewer_pairs(self, ranked):
        assert len(select_tiled(ranked, 8, 0)) >= len(select_tiled(ranked, 8, 200))

    def test_returned_in_positional_order_not_score_order(self, ranked):
        offsets = [c.offset for c in select_tiled(ranked, 5)]
        assert offsets == sorted(offsets), "an order sheet reads positionally"

    def test_capacity_limits_are_not_an_error(self, ranked):
        """Asking for more than fits returns what fits, for the caller to notice."""
        got = select_tiled(ranked, 500)
        assert 1 < len(got) < 500

    def test_best_candidate_is_always_included(self, ranked):
        assert ranked[0] in select_tiled(ranked, 4)

    def test_empty_input_gives_nothing(self):
        assert select_tiled([], 3) == []

    def test_bad_arguments_are_rejected(self, ranked):
        with pytest.raises(ValueError, match="n_pairs"):
            select_tiled(ranked, 0)
        with pytest.raises(ValueError, match="min_spacing"):
            select_tiled(ranked, 2, -1)


class TestNaming:
    def test_a_single_pair_gets_the_unsuffixed_names(self):
        """The three verified designs must keep regenerating byte-identically."""
        assert tiled_names("Gene1", "prom", 0, 1) == ("P1L_Gene1_prom", "P1R_Gene1_prom")

    def test_tiled_names_are_numbered_from_one_and_padded(self):
        assert tiled_names("G", "p", 0, 12)[0] == "P1L_G_p_t01"
        assert tiled_names("G", "p", 9, 12)[0] == "P1L_G_p_t10"

    def test_every_name_in_a_tile_is_distinct(self, sys_a, ranked):
        pairs = build_tiled_pairs("prom", "Gene1", sys_a, select_tiled(ranked, 6))
        names = [n for p in pairs for n in (p.p1l_name, p.p1r_name)]
        assert len(set(names)) == len(names) == 12


class TestDesignResultShape:
    def test_constructing_without_pairs_still_works(self, sys_a, region):
        """Everything predating tiling constructs DesignResult this way."""
        seq, iv = region
        cands = scan(seq, iv, top=3)
        pair = build_pair("prom", sys_a, cands[0], "P1L", "P1R")
        res = DesignResult(Locus("prom", "G", "", iv, sys_a.name), sys_a, seq,
                           cands, pair, [], not_screened("x"))
        assert res.pairs == [pair]
        assert res.n_pairs == 1 and not res.tiled
        assert res.pair_screens == [res.specificity]

    def test_rows_pads_missing_screens_explicitly(self, sys_a, region):
        """A hand-built result must not borrow pair 1's screen for pair 2."""
        seq, iv = region
        cands = scan(seq, iv, top=len(seq))
        pairs = build_tiled_pairs("prom", "G", sys_a, select_tiled(cands, 3))
        res = DesignResult(Locus("prom", "G", "", iv, sys_a.name), sys_a, seq,
                           cands[:5], pairs[0], [], not_screened("first"),
                           pairs=pairs, pair_screens=[not_screened("first")])
        rows = list(res.rows())
        assert len(rows) == 3
        assert "no screen recorded" in rows[1][1].reason
        assert not rows[1][1].screened


class TestEveryPairIsScreened:
    """Ordering nine unchecked pairs alongside one checked one is worse than
    not tiling: they all look equally validated on the sheet."""

    def test_one_screen_per_pair(self, sys_a, region):
        # design_locus fetches from UCSC, so the assembled result is exercised
        # here and the fetch path is covered by the end-to-end CLI run.
        built = make_result(sys_a, region, 5)
        assert len(built.pair_screens) == built.n_pairs == 5
        assert all(s is not None for s in built.pair_screens)

    def test_design_locus_screens_each_selected_pair(self, sys_a, region,
                                                     monkeypatch):
        """The real path, with the network and RepeatMasker stubbed out."""
        from mypibypy import pipeline
        seq, iv = region
        monkeypatch.setattr(pipeline, "fetch_sequence", lambda interval: seq)
        monkeypatch.setattr(pipeline, "rmsk_hits", lambda interval: [])
        monkeypatch.setattr(pipeline, "mask_repeats", lambda s, i, h: s)
        calls = []

        def fake_screen(left, right, backend=None, **kw):
            calls.append((left, right))
            return not_screened("stub")

        monkeypatch.setattr(pipeline, "screen_pair", fake_screen)
        # Locus carries the registry KEY, not AmplifierSystem.name.
        locus = Locus("prom", "Gene1", "synthetic", iv, sorted(systems())[0])
        res = pipeline.design_locus(locus, n_pairs=4)
        assert res.n_pairs == 4
        assert len(calls) == 4, "every selected pair must be screened, not just the first"
        assert len({c[0] for c in calls}) == 4, "and each with its own footprint"

    def test_rows_yields_a_screen_for_every_pair(self, sys_a, region):
        built = make_result(sys_a, region, 4)
        assert len({id(s) for _, s, _ in built.rows()}) == 4


class TestFusionRefusesTiling:
    """A junction is one position -- the design IS the gap on the breakpoint."""

    @pytest.fixture
    def junction(self):
        from mypibypy.loci import load_fusions
        from pathlib import Path
        import mypibypy.loci as loci_mod
        return load_fusions(Path(loci_mod.__file__).parent / "loci.example.json")[0]

    def test_more_than_one_pair_is_refused_not_ignored(self, junction):
        with pytest.raises(FusionIsSinglePair) as exc:
            design_fusion(junction, screen=False, n_pairs=2)
        message = str(exc.value)
        assert "breakpoint" in message
        assert "wholly inside one partner" in message, (
            "the message must say why, not just that it refuses")

    def test_one_pair_still_yields_sense_and_antisense(self, junction):
        """Two results is the strand axis, not tiling."""
        results = design_fusion(junction, screen=False, n_pairs=1)
        assert len(results) == 2
        assert all(r.n_pairs == 1 for r in results)

    def test_the_default_is_single(self, junction):
        assert len(design_fusion(junction, screen=False)) == 2


class TestWorkbookWithManyPairs:
    def _cells(self, tmp_path, results, name="out.xlsx"):
        path = tmp_path / name
        write_workbook(str(path), results, "test note")
        wb = load_workbook(path)
        return wb

    def test_order_sheet_lists_every_primary_under_a_distinct_name(
            self, tmp_path, sys_a, region):
        wb = self._cells(tmp_path, [make_result(sys_a, region, 6)])
        names = [r[0] for r in wb["ORDER_ssDNA"].iter_rows(
            min_row=5, max_col=1, values_only=True) if r[0]]
        assert len(names) == 12
        assert len(set(names)) == 12

    def test_reused_parts_do_not_grow_with_pair_count(self, tmp_path, sys_a, region):
        """Tiling into an existing pool is the point: P2/P3/P4 stay as they are."""
        def counts(n):
            wb = self._cells(tmp_path, [make_result(sys_a, region, n)], f"n{n}.xlsx")
            rows = [r for r in wb["Parts list"].iter_rows(
                min_row=4, max_col=5, values_only=True) if r[3]]
            return (sum(1 for r in rows if r[3] == "NEW"),
                    sum(1 for r in rows if r[3] == "REUSED"))
        new1, reused1 = counts(1)
        new5, reused5 = counts(5)
        assert (new1, new5) == (2, 10)
        assert reused1 == reused5, "reused amplifier count must not scale with pairs"

    def test_tiled_locus_sheet_has_a_row_per_pair(self, tmp_path, sys_a, region):
        wb = self._cells(tmp_path, [make_result(sys_a, region, 4)])
        ws = wb["Gene1_prom"]
        header = [c for c in next(ws.iter_rows(min_row=4, max_col=14, values_only=True))]
        assert header[0] == "#" and "Genome screen" in header
        numbers = [r[0] for r in ws.iter_rows(min_row=5, max_col=1, values_only=True)
                   if isinstance(r[0], int)]
        assert numbers == [1, 2, 3, 4]

    def test_notes_count_all_the_oligos(self, tmp_path, sys_a, region):
        wb = self._cells(tmp_path, [make_result(sys_a, region, 3)])
        text = "\n".join(str(c) for row in wb["DESIGN_NOTES"].iter_rows(values_only=True)
                         for c in row if c)
        assert "6 ssDNA P1 primaries" in text
        assert "tiled" in text


class TestSinglePairIsUnchanged:
    """The regression guard. Single-pair output must not drift."""

    def test_names_carry_no_tile_suffix(self, tmp_path, sys_a, region):
        wb = TestWorkbookWithManyPairs()._cells(
            tmp_path, [make_result(sys_a, region, 1)], "single.xlsx")
        names = [r[0] for r in wb["ORDER_ssDNA"].iter_rows(
            min_row=5, max_col=1, values_only=True) if r[0]]
        assert names == ["P1L_Gene1_prom", "P1R_Gene1_prom"]
        assert not any("_t0" in n for n in names)

    def test_single_pair_keeps_the_form_layout_not_the_table(
            self, tmp_path, sys_a, region):
        """One pair reads as a form; many read as a table. Don't merge them."""
        wb = TestWorkbookWithManyPairs()._cells(
            tmp_path, [make_result(sys_a, region, 1)], "single2.xlsx")
        ws = wb["Gene1_prom"]
        labels = [r[0] for r in ws.iter_rows(max_col=1, values_only=True) if r[0]]
        assert "Chosen footprint" in labels
        assert "#" not in labels, "the single-pair sheet must not become a table"

    def test_notes_do_not_mention_tiling_for_one_pair(self, tmp_path, sys_a, region):
        wb = TestWorkbookWithManyPairs()._cells(
            tmp_path, [make_result(sys_a, region, 1)], "single3.xlsx")
        text = "\n".join(str(c) for row in wb["DESIGN_NOTES"].iter_rows(values_only=True)
                         for c in row if c)
        assert "2 ssDNA P1 primaries" in text
        assert "tiled" not in text

    def test_a_tiled_design_of_one_equals_an_untiled_one(self, sys_a, region):
        """n_pairs=1 must go down the identical path, not a parallel one."""
        seq, iv = region
        cands = scan(seq, iv, top=len(seq))
        tiled = build_tiled_pairs("prom", "Gene1", sys_a, select_tiled(cands, 1))
        direct = build_pair("prom", sys_a, cands[0],
                            "P1L_Gene1_prom", "P1R_Gene1_prom")
        assert len(tiled) == 1
        assert (tiled[0].p1l, tiled[0].p1r) == (direct.p1l, direct.p1r)
        assert (tiled[0].p1l_name, tiled[0].p1r_name) == (direct.p1l_name,
                                                          direct.p1r_name)


# --- re-packing when greedy falls short -------------------------------------


def _fits(chosen, spacing):
    offs = sorted(c.offset for c in chosen)
    return all(b - a >= FOOTPRINT + spacing for a, b in zip(offs, offs[1:]))


def _windows(seed, length=260):
    rng = random.Random(seed)
    seq = "".join(rng.choice("ACGT") for _ in range(length))
    return scan(seq, GenomicInterval("hg38", "chr1", 1, length), top=length)


@pytest.mark.parametrize("seed", range(6))
@pytest.mark.parametrize("n_pairs,spacing", [(2, 10), (3, 20), (4, 5), (5, 0)])
def test_repack_matches_brute_force(seed, n_pairs, spacing):
    """The most pairs that fit, and the least total penalty among those."""
    windows = _windows(seed)[::7]           # a small set, so brute force is cheap
    best_k, best_cost = 0, 0.0
    for k in range(n_pairs, 0, -1):
        costs = [sum(c.penalty for c in combo) for combo in combinations(windows, k)
                 if _fits(combo, spacing)]
        if costs:
            best_k, best_cost = k, min(costs)
            break
    got = repack(windows, n_pairs, spacing)
    assert len(got) == best_k and _fits(got, spacing)
    assert sum(c.penalty for c in got) == pytest.approx(best_cost)


@pytest.mark.parametrize("seed", range(6))
def test_greedy_is_kept_whenever_it_places_every_pair(seed):
    windows = _windows(seed)
    for n_pairs, spacing in [(1, 3), (2, 3), (3, 10)]:
        greedy = select_tiled(windows, n_pairs, spacing)
        if len(greedy) == n_pairs:
            assert tile(windows, n_pairs, spacing) == (greedy, 0)


def test_a_shortfall_is_re_packed_and_reported():
    """A request greedy cannot meet but a different layout can."""
    for seed in range(40):
        windows = _windows(seed)
        greedy = select_tiled(windows, 4, 15)
        if len(greedy) < 4:
            packed, from_greedy = tile(windows, 4, 15)
            if len(packed) > len(greedy):
                assert from_greedy == len(greedy) and _fits(packed, 15)
                return
    pytest.fail("no seed where greedy fell short and re-packing helped")


def test_repack_honours_a_cost():
    windows = _windows(1)[::5]
    cost = lambda c: -c.offset                   # noqa: E731 -- prefer late windows
    got = repack(windows, 2, 0, cost)
    assert got[-1].offset == max(c.offset for c in windows)
