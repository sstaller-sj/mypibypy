"""Render a piFISH order workbook from design results.

Follows the lab's established order-sheet conventions: an ORDER_ssDNA sheet in
IDT name/sequence format, a Parts list with =LEN() columns, and a per-locus
design sheet documenting provenance. Only the new P1 primaries are flagged for
ordering; the P2/P3/P4 amplifier oligos are reused from the existing pool.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .expression import EXPRESSION_IS_UNCALIBRATED
from .pipeline import DesignResult

FONT = "Arial"
H = Font(name=FONT, bold=True, size=11)
B = Font(name=FONT, size=11)
MONO = Font(name=FONT, size=10)
TITLE = Font(name=FONT, bold=True, size=13)
NEW_FILL = PatternFill("solid", fgColor="FFF2CC")     # pale yellow = order this
REUSE_FILL = PatternFill("solid", fgColor="E2EFDA")   # pale green = already in hand
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")     # pale orange = unverified


#: Characters Excel refuses in a worksheet title, plus its 31-char limit.
_BAD_SHEET_CHARS = str.maketrans({c: "-" for c in ":\\/?*[]"})
_SHEET_TITLE_MAX = 31


def _sheet_title(wb, name: str) -> str:
    """A worksheet title Excel will accept, unique within `wb`.

    Fusion genes are written ``GENEA::GENEB`` and a colon is illegal in a sheet
    title, so this is not hypothetical -- and a 31-character ceiling makes
    collisions between long names likely enough to handle.
    """
    base = (name.translate(_BAD_SHEET_CHARS).strip() or "sheet")[:_SHEET_TITLE_MAX]
    title, n = base, 2
    while title in wb.sheetnames:
        suffix = f"_{n}"
        title = base[: _SHEET_TITLE_MAX - len(suffix)] + suffix
        n += 1
    return title


def _specificity_text(rep) -> str:
    """One cell describing the genome screen, including it not having happened.

    This function is the whole reason `SpecificityReport` has a `screened` flag.
    An earlier workbook asserted a specificity result unconditionally and put an
    unverified "repeat-free single-copy" claim onto a real order sheet.
    """
    if not rep.screened:
        return ("NO GENOME-WIDE SPECIFICITY SCREEN WAS RUN for this locus "
                f"({rep.reason}). Nothing here establishes single-copy status. "
                "Configure a local bowtie v1 index ($MYPIBYPY_BOWTIE_INDEX) or "
                "BLAST database ($MYPIBYPY_BLAST_DB) and re-run before ordering.")
    verdict = {
        "pass": "PASS -- no off-target site where both halves land together",
        "review": "REVIEW -- off-target co-localization found, see below",
        "fail": "FAIL -- both halves land together elsewhere; do not order",
    }.get(rep.verdict, rep.verdict.upper())
    text = f"{verdict}. {rep.summary()}"
    for site in rep.pair_sites[:5]:
        text += f"\n  - {site}"
    if len(rep.pair_sites) > 5:
        text += f"\n  - ... and {len(rep.pair_sites) - 5} more"
    if rep.expression_evaluated:
        text += f"\n{EXPRESSION_IS_UNCALIBRATED}"
    return text


def _folding_text(fold) -> str:
    """One cell for probe self-folding, including it not having been run."""
    return fold.summary()


def _folding_flagged(fold) -> bool:
    """Worth a reader's eye -- a word for a colour, not a reject."""
    return fold.evaluated and fold.limiting.label == "folded"


def _accessibility_text(res) -> str:
    """One cell for target-site accessibility, including it not being applicable.

    Same rule as the specificity cell: a number that was never computed must not
    render as a good one. The uncalibrated caveat travels with the number,
    because a reader looking at an order sheet is exactly who might otherwise
    treat it as a threshold.
    """
    if not res.evaluated:
        return f"NOT EVALUATED -- {res.reason}"
    limiting = res.limiting_pad
    return (
        f"{limiting.label.upper()} -- limiting pad is the {limiting.pad} at "
        f"P={res.pair_nucleation_p:.3f} for a {res.nucleation_nt}-nt unpaired run "
        f"(left {res.left.best_nucleation_p:.3f} / right "
        f"{res.right.best_nucleation_p:.3f}; mean per-base unpaired "
        f"{res.left.mean_unpaired:.2f} / {res.right.mean_unpaired:.2f}). "
        f"{res.tool}. {res.caveat}"
    )


def _hdr(ws, row, values, widths=None):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = H
        c.fill = HDR_FILL
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def _notes_sheet(wb, results, source_note):
    ws = wb.create_sheet("DESIGN_NOTES")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    # Every field below is derived from the run that produced this workbook.
    # Earlier versions hardcoded the gene, the oligo count and the reused
    # amplifiers, which then disagreed with any design but the one they were
    # written for.
    genes = sorted({r.locus.gene for r in results})
    builds = sorted({r.locus.interval.build for r in results})
    reused = []
    for res in results:
        s = res.system
        label = f"P2/P3/{s.p4_name} ({s.fluorophore})"
        if label not in reused:
            reused.append(label)
    rows = [
        ("piFISH split-probe design", None, TITLE),
        ("Gene", ", ".join(genes), None),
        ("Genome build", ", ".join(builds), None),
        ("Method", "piFISH split primary (P1L/P1R) into reused P2/P3/P4 amplifier pool", None),
        ("Footprint", "25-nt left pad + 1-nt gap + 25-nt right pad = 51 nt; probes hybridize to the plus strand", None),
        ("P1L", "arm_left + revcomp(left25)", None),
        ("P1R", "revcomp(right25) + arm_right", None),
        ("New oligos to order",
         f"{sum(2 * r.n_pairs for r in results)} ssDNA P1 primaries "
         f"(see ORDER_ssDNA)"
         + (f"; {len(results)} locus/loci tiled with "
            f"{', '.join(str(r.n_pairs) for r in results)} pair(s) each"
            if any(r.tiled for r in results) else ""), None),
        ("Reused (in hand)", "; ".join(reused), None),
        ("", "", None),
        ("Loci", "", H),
    ]
    r = 1
    for a, b, font in rows:
        ca = ws.cell(row=r, column=1, value=a); ca.font = font or H
        cb = ws.cell(row=r, column=2, value=b); cb.font = B
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    for res in results:
        L = res.locus
        ws.cell(row=r, column=1, value=f"{L.gene} {L.name}").font = H
        ws.cell(row=r, column=2, value=(
            f"{L.description}; {L.interval}; -> {res.system.fluorophore} "
            f"[system {res.system.name}]; footprint {res.best.interval}; "
            f"GC L/R {res.best.left.gc*100:.0f}%/{res.best.right.gc*100:.0f}%, "
            f"Tm L/R {res.best.left.tm:.1f}/{res.best.right.tm:.1f} C; "
            f"RepeatMasker hits: {len(res.rmsk)}"
        )).font = B
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Specificity").font = H
    # Report what actually ran, per locus. A workbook must never assert a screen
    # that was not performed: RepeatMasker only says a footprint avoids
    # annotated repeats, which is not the same as being single-copy genome-wide,
    # and a pair-aware screen is the only thing that speaks to piFISH signal.
    total_rmsk = sum(len(res.rmsk) for res in results)
    n = len(results)
    ws.cell(row=r, column=2, value=(
        f"RepeatMasker: {total_rmsk} annotated repeat(s) overlapping the "
        f"{n} chosen footprint{'s' if n != 1 else ''}"
        + ("; none placed in annotated repetitive DNA. " if total_rmsk == 0
           else " -- review before ordering. ")
        + "RepeatMasker does not establish single-copy status; the genome-wide "
          "result per locus is below."
    )).font = B
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
    for res in results:
        L = res.locus
        ws.cell(row=r, column=1, value=f"  {L.gene} {L.name}").font = B
        cell = ws.cell(row=r, column=2, value=_specificity_text(res.specificity))
        cell.font = B
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if not res.specificity.screened or res.specificity.verdict == "fail":
            cell.fill = WARN_FILL
        r += 1
    if any(res.accessibility.evaluated for res in results):
        ws.cell(row=r, column=1, value="Accessibility").font = H
        ws.cell(row=r, column=2, value=(
            "Target-site secondary structure on the transcript (RNA targets "
            "only; DNA targets are denatured by the protocol). Per locus below."
        )).font = B
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
        for res in results:
            L = res.locus
            ws.cell(row=r, column=1, value=f"  {L.gene} {L.name}").font = B
            cell = ws.cell(row=r, column=2, value=_accessibility_text(res.accessibility))
            cell.font = B
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if res.accessibility.evaluated and res.accessibility.pair_nucleation_p < 0.2:
                cell.fill = WARN_FILL
            r += 1
        r += 1

    if any(res.folding_at(i).evaluated for res in results for i in range(res.n_pairs)):
        ws.cell(row=r, column=1, value="Probe folding").font = H
        ws.cell(row=r, column=2, value=(
            "Self-folding of each synthesized P1L/P1R (RNAfold partition function, "
            "DNA parameters, primary-hyb conditions): mean probability that a pad "
            "base and an arm base is paired, with the MFE as context. The more "
            "folded primary per pair is shown. At 6x SSC random sequence is ~40% "
            "paired, so 'open' runs to 50%; the words are not thresholds."
        )).font = B
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
        for res in results:
            L = res.locus
            for i in range(res.n_pairs):
                fold = res.folding_at(i)
                label = f"  {L.gene} {L.name}" + (f" #{i + 1}" if res.tiled else "")
                ws.cell(row=r, column=1, value=label).font = B
                cell = ws.cell(row=r, column=2, value=(
                    str(fold.limiting) if fold.evaluated
                    else f"NOT EVALUATED -- {fold.reason}"))
                cell.font = B
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if _folding_flagged(fold):
                    cell.fill = WARN_FILL
                r += 1
        r += 1

    ws.cell(row=r, column=1, value="Source/provenance").font = H
    ws.cell(row=r, column=2, value=source_note).font = B
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def _order_sheet(wb, results):
    ws = wb.create_sheet("ORDER_ssDNA")
    ws.cell(row=1, column=1, value="IDT order - NEW ssDNA oligos (P1 split primaries only)").font = TITLE
    ws.cell(row=2, column=1, value="100 nmole DNA oligo, standard desalting").font = B
    _hdr(ws, 4, ["NAME", "SEQUENCE", "LENGTH"], widths=[22, 60, 10])
    r = 5
    for res in results:
        for pair in res.pairs:
            for name, seq in ((pair.p1l_name, pair.p1l), (pair.p1r_name, pair.p1r)):
                ws.cell(row=r, column=1, value=name).font = B
                cs = ws.cell(row=r, column=2, value=seq); cs.font = MONO
                cl = ws.cell(row=r, column=3, value=f"=LEN(B{r})"); cl.font = B
                for col in (1, 2, 3):
                    ws.cell(row=r, column=col).fill = NEW_FILL
                r += 1
    ws.freeze_panes = "A5"
    return ws


def _parts_sheet(wb, results):
    ws = wb.create_sheet("Parts list")
    ws.cell(row=1, column=1, value="Full parts list (NEW = order; REUSED = existing pool)").font = TITLE
    _hdr(ws, 3, ["Name", "Sequence (5'->3')", "LENGTH", "Status", "Channel"], widths=[22, 60, 10, 10, 10])
    r = 4
    for res in results:
        ch = res.system.fluorophore
        # Every tiled pair is a NEW oligo to order. The P2/P3/P4 amplifiers below
        # stay REUSED however many pairs there are -- that is the whole point of
        # tiling into an existing pool, and the count of reused parts must not
        # grow with the number of pairs.
        for pair in res.pairs:
            for name, seq in ((pair.p1l_name, pair.p1l), (pair.p1r_name, pair.p1r)):
                _part_row(ws, r, name, seq, "NEW", ch, NEW_FILL); r += 1
    seen = set()
    for res in results:
        s = res.system
        for name, seq in ((f"P2_{s.name}", s.p2), (f"P3_{s.name}", s.p3), (s.p4_name, s.p4)):
            if name in seen:
                continue
            seen.add(name)
            _part_row(ws, r, name, seq, "REUSED", s.fluorophore, REUSE_FILL); r += 1
    ws.freeze_panes = "A4"
    return ws


def _part_row(ws, r, name, seq, status, channel, fill):
    ws.cell(row=r, column=1, value=name).font = B
    ws.cell(row=r, column=2, value=seq).font = MONO
    ws.cell(row=r, column=3, value=f"=LEN(B{r})").font = B
    ws.cell(row=r, column=4, value=status).font = B
    ws.cell(row=r, column=5, value=channel).font = B
    for col in range(1, 6):
        ws.cell(row=r, column=col).fill = fill


def _tiled_locus_sheet(wb, res):
    """Per-locus sheet for a tiled design: one row per ordered pair.

    A separate layout rather than a generalised one, so the single-pair sheet
    stays byte-identical -- the three verified designs must keep regenerating
    exactly. Reads as a table because that is what someone checking ten pairs
    wants; the single-pair sheet reads as a form because that is what someone
    checking one wants.
    """
    L, sysm = res.locus, res.system
    ws = wb.create_sheet(_sheet_title(wb, f"{L.gene}_{L.name}"))
    ws.cell(row=1, column=1, value=f"{L.gene} {L.name} - {L.description}").font = TITLE
    ws.cell(row=2, column=1, value=(
        f"{res.n_pairs} tiled pairs into {sysm.fluorophore} (system {sysm.name}); "
        f"target region {L.interval}; {2 * res.n_pairs} new ssDNA oligos. "
        f"Each pair is screened independently -- see the per-pair columns."
    )).font = B
    ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    headers = ["#", "Footprint", "left25 (plus)", "gap", "right25 (plus)",
               "GC L/R", "Tm L/R", "QC flags", "P1L name", "P1L sequence",
               "P1R name", "P1R sequence", "Genome screen", "Accessibility",
               "Probe folding"]
    _hdr(ws, 4, headers,
         widths=[4, 26, 27, 5, 27, 11, 13, 24, 22, 52, 22, 52, 46, 46, 46])
    r = 5
    for i, (pair, screen, access) in enumerate(res.rows(), start=1):
        cand = pair.candidate
        values = [
            i, str(cand.interval), cand.left.seq, cand.gap_base, cand.right.seq,
            f"{cand.left.gc*100:.0f}%/{cand.right.gc*100:.0f}%",
            f"{cand.left.tm:.1f}/{cand.right.tm:.1f}",
            ", ".join(cand.flags) or "none",
            pair.p1l_name, pair.p1l, pair.p1r_name, pair.p1r,
            _specificity_text(screen), _accessibility_text(access),
            _folding_text(res.folding_at(i - 1)),
        ]
        for j, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = MONO if j in (3, 5, 10, 12) else B
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        # Flag the pairs a reader must look at, rather than leaving them to
        # find the one bad row in ten.
        if not screen.screened or screen.verdict == "fail":
            ws.cell(row=r, column=13).fill = WARN_FILL
        if access.evaluated and access.pair_nucleation_p < 0.2:
            ws.cell(row=r, column=14).fill = WARN_FILL
        if _folding_flagged(res.folding_at(i - 1)):
            ws.cell(row=r, column=15).fill = WARN_FILL
        r += 1
    ws.freeze_panes = "B5"
    r += 1
    ws.cell(row=r, column=1, value="Region sequence (plus, 5'->3')").font = H
    ws.cell(row=r, column=2, value=res.region_seq).font = MONO
    return ws


def _locus_sheet(wb, res):
    if res.tiled:
        return _tiled_locus_sheet(wb, res)
    L, c, p, s = res.locus, res.best, res.pair, res.system
    ws = wb.create_sheet(_sheet_title(wb, f"{L.gene}_{L.name}"))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80
    def kv(r, k, v, mono=False):
        ws.cell(row=r, column=1, value=k).font = H
        cell = ws.cell(row=r, column=2, value=v)
        cell.font = MONO if mono else B
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=1, value=f"{L.gene} {L.name} - {L.description}").font = TITLE
    kv(3, "Channel", f"{s.fluorophore} (system {s.name})")
    kv(4, "Target region", str(L.interval))
    kv(5, "Region length", f"{L.interval.length} nt")
    kv(6, "Chosen footprint", str(c.interval))
    kv(7, "left25 (plus)", c.left.seq, mono=True)
    kv(8, "gap (skipped)", c.gap_base, mono=True)
    kv(9, "right25 (plus)", c.right.seq, mono=True)
    kv(10, "left GC / Tm", f"{c.left.gc*100:.0f}% / {c.left.tm:.1f} C")
    kv(11, "right GC / Tm", f"{c.right.gc*100:.0f}% / {c.right.tm:.1f} C")
    kv(12, "QC flags", ", ".join(c.flags) or "none")
    kv(13, "Genome screen", _specificity_text(res.specificity))
    kv(14, "Target-site accessibility", _accessibility_text(res.accessibility))
    kv(15, "Probe self-folding", _folding_text(res.folding))
    kv(16, p.p1l_name, p.p1l, mono=True)
    kv(17, p.p1r_name, p.p1r, mono=True)
    kv(19, "Region sequence (plus, 5'->3')", res.region_seq, mono=True)
    return ws


def write_workbook(path: str, results: list[DesignResult], source_note: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    _notes_sheet(wb, results, source_note)
    _order_sheet(wb, results)
    _parts_sheet(wb, results)
    for res in results:
        _locus_sheet(wb, res)
    wb.save(path)
    return path
