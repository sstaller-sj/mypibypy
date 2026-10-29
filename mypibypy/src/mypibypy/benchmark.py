"""Where the lab's ordered designs rank under a given scoring.

The designs the lab has already ordered (and, for some, imaged) are the only
ground truth this program has. Before the ranking changes, this measures where
each ordered pair sits under the current scoring ("v1") and, once it exists,
under the new one -- so a change that pushes a working design down the list is
seen before it becomes the default, not after.

Input is the lab's order workbooks. Each design sheet carries its target
region, its sequence and the chosen footprint(s), in one of two layouts:

* a single pair: label/value rows -- ``Target region``, ``Chosen footprint``,
  ``left25 (plus)``, ``right25 (plus)``, ``Region sequence (plus, 5'->3')``;
* tiled pairs: a table headed ``# | Footprint | left25 (plus) | gap |
  right25 (plus) | ...``, with the region named on the sheet's second row and
  ``Region sequence`` below the table.

Fusion sheets have neither and are skipped: a junction is one fixed footprint,
so there is no ranking choice to test.

NOTHING HERE PRINTS A SEQUENCE. Ordered sequences are the lab's and
unpublished; reports carry labels, positions, ranks and costs only.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .design import DEFAULT_MIN_SPACING_NT, FOOTPRINT, LEFT_LEN, GAP_LEN, tile
from .sequence import GenomicInterval

#: "mm10 chr7:46,371,076-46,371,126" or "ENST00000458748 transcript:2-52".
_PLACE_RE = re.compile(r"(?P<build>\S+)\s+(?P<chrom>[\w.]+):(?P<start>[\d,]+)-(?P<end>[\d,]+)")
_REGION_IN_TEXT = re.compile(r"target region\s+(?P<place>\S+\s+[\w.]+:[\d,]+-[\d,]+)",
                             re.IGNORECASE)


@dataclass(frozen=True)
class OrderedPair:
    """One ordered pair, by where its footprint sits in the region."""

    label: str
    offset: int                                   # 0-based within the region
    left25: str = field(default="", repr=False)
    right25: str = field(default="", repr=False)


@dataclass
class OrderedDesign:
    workbook: str
    sheet: str
    kind: str                                     # "genomic" | "rna"
    region: str                                   # e.g. "mm10 chr7:46,371,022-46,371,467"
    seq: str = field(repr=False)
    interval: GenomicInterval | None = None       # genomic designs
    accession: str = ""                           # RNA designs
    pairs: list[OrderedPair] = field(default_factory=list)
    #: Anything that did not add up -- a pad that is not where its footprint
    #: says. Reported, never silently corrected.
    problems: list[str] = field(default_factory=list)


def _place(text: str) -> tuple[str, str, int, int] | None:
    m = _PLACE_RE.search(str(text or ""))
    if not m:
        return None
    return (m["build"], m["chrom"], int(m["start"].replace(",", "")),
            int(m["end"].replace(",", "")))


def _clean(seq) -> str:
    return re.sub(r"[^ACGTUacgtu]", "", str(seq or "")).upper().replace("U", "T")


def _check_pads(design: OrderedDesign) -> None:
    """Each ordered pad must be where its footprint says, in the region given."""
    for pair in design.pairs:
        got_l = design.seq[pair.offset:pair.offset + LEFT_LEN]
        start_r = pair.offset + LEFT_LEN + GAP_LEN
        got_r = design.seq[start_r:start_r + LEFT_LEN]
        if pair.left25 and got_l != _clean(pair.left25):
            design.problems.append(f"{pair.label}: left pad is not at its footprint")
        if pair.right25 and got_r != _clean(pair.right25):
            design.problems.append(f"{pair.label}: right pad is not at its footprint")


def _design(path: Path, sheet: str, region_text: str, seq: str,
            pairs: list[tuple[str, str, str, str]]) -> OrderedDesign | None:
    """Build a design from a region and (label, footprint, left25, right25) rows."""
    region = _place(region_text)
    first = _place(pairs[0][1]) if pairs else None
    if first is None:
        return None
    rna = first[1] == "transcript"
    if rna:
        accession, base = first[0], 1         # transcript coordinates are 1-based
        interval = None
    else:
        if region is None:
            return None
        accession, base = "", region[2]
        interval = GenomicInterval(region[0], region[1], region[2], region[3])
    ordered = []
    for label, footprint, left, right in pairs:
        place = _place(footprint)
        if place is None:
            continue
        ordered.append(OrderedPair(label, place[2] - base, left, right))
    design = OrderedDesign(path.name, sheet, "rna" if rna else "genomic",
                           region_text or f"{accession} transcript", seq,
                           interval=interval, accession=accession, pairs=ordered)
    if interval is not None and interval.length != len(seq):
        design.problems.append(f"region is {interval.length} nt but its sequence "
                               f"is {len(seq)} nt")
    _check_pads(design)
    return design


def read_workbook(path: str | Path) -> list[OrderedDesign]:
    """Every design sheet in an order workbook that has a footprint to rank."""
    from openpyxl import load_workbook

    path = Path(path)
    book = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in book.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        labels = {str(r[0]).strip(): (r[1] if len(r) > 1 else None)
                  for r in rows if r and r[0] is not None}
        seq = _clean(next((v for k, v in labels.items()
                           if k.lower().startswith("region sequence")), ""))
        header = next((i for i, r in enumerate(rows)
                       if r and str(r[0]).strip() == "#" and len(r) > 1
                       and str(r[1]).strip().lower() == "footprint"), None)
        if header is not None:
            names = [str(c or "").strip().lower() for c in rows[header]]
            col = {n: i for i, n in enumerate(names)}
            pairs = []
            for r in rows[header + 1:]:
                if not r or r[0] is None or not str(r[0]).strip().isdigit():
                    break
                get = lambda n: r[col[n]] if n in col and col[n] < len(r) else ""  # noqa: E731
                pairs.append((f"{ws.title} #{str(r[0]).strip()}", str(get("footprint")),
                              str(get("left25 (plus)") or ""), str(get("right25 (plus)") or "")))
            region_text = next((m["place"] for r in rows[:3] for c in r if c
                                for m in [_REGION_IN_TEXT.search(str(c))] if m), "")
            design = _design(path, ws.title, region_text, seq, pairs)
        elif "Target region" in labels and "Chosen footprint" in labels:
            design = _design(path, ws.title, str(labels["Target region"]), seq,
                             [(ws.title, str(labels["Chosen footprint"]),
                               str(labels.get("left25 (plus)") or ""),
                               str(labels.get("right25 (plus)") or ""))])
        else:
            continue
        if design is not None and design.pairs and design.seq:
            out.append(design)
    return out


# --- ranking ------------------------------------------------------------------

def rank_v1(design: OrderedDesign, rmsk: list[dict] | None = None):
    """(windows in the order the pipeline tiles from, cost function) -- today's
    scoring, through the pipeline's own ranking functions."""
    from .loci import RnaTarget
    from .pipeline import rank_locus_windows, rank_rna_windows

    if design.kind == "genomic":
        return rank_locus_windows(design.seq, design.interval, rmsk or []), None
    ranking = rank_rna_windows(design.seq, RnaTarget(design.sheet, "GENE",
                                                     design.accession, ""))
    return ranking.ranked, (lambda c: ranking.costs[c.offset])


#: Scorings the benchmark can compare. v2 joins this when it exists.
RANKERS = {"v1": rank_v1}


@dataclass(frozen=True)
class PairRank:
    """Where one ordered pair sits under one scoring. No sequence."""

    scoring: str
    workbook: str
    design: str
    pair: str
    footprint: str            # 1-based within the region, e.g. "55-105"
    rank: int | None          # 1 = first the tiler would take; None = not a candidate
    windows: int
    cost: float | None
    best_cost: float
    picked: bool              # chosen when the design's pair count is tiled
    note: str = ""


def _spacing(design: OrderedDesign) -> int:
    """The default spacing, or less if the ordered pairs sit closer than it."""
    offsets = sorted(p.offset for p in design.pairs)
    gaps = [b - a - FOOTPRINT for a, b in zip(offsets, offsets[1:])]
    return max(0, min([DEFAULT_MIN_SPACING_NT, *gaps]))


def rank_design(design: OrderedDesign, scoring: str = "v1",
                rmsk: list[dict] | None = None) -> list[PairRank]:
    ranked, cost = RANKERS[scoring](design, rmsk)
    cost = cost or (lambda c: c.penalty)
    position = {c.offset: i for i, c in enumerate(ranked, 1)}
    by_offset = {c.offset: c for c in ranked}
    chosen, _ = tile(ranked, len(design.pairs), _spacing(design), cost)
    picked = {c.offset for c in chosen}
    best = min((cost(c) for c in ranked), default=float("nan"))
    out = []
    for pair in design.pairs:
        cand = by_offset.get(pair.offset)
        out.append(PairRank(
            scoring, design.workbook, design.sheet, pair.label,
            f"{pair.offset + 1}-{pair.offset + FOOTPRINT}",
            position.get(pair.offset), len(ranked),
            round(cost(cand), 3) if cand else None, round(best, 3),
            pair.offset in picked,
            "" if cand else "not a candidate (repeat-masked or invalid window)"))
    return out


def markdown(rows: list[PairRank], outcomes: dict[str, str] | None = None,
             names: bool = False) -> str:
    """The benchmark as a table.

    `outcomes` maps a pair's label (its sheet name, or "sheet #n" for a tiled
    pair) to what imaging showed -- "worked", "failed"; anything missing is
    "unknown". Without `names`, workbooks, designs and pairs are shown by index
    only, as the repo's other validation scripts do, since sheet names carry
    the lab's target genes.
    """
    outcomes = outcomes or {}
    books = list(dict.fromkeys(r.workbook for r in rows))
    designs = list(dict.fromkeys((r.workbook, r.design) for r in rows))
    pairs = list(dict.fromkeys((r.workbook, r.design, r.pair) for r in rows))
    lines = ["| scoring | design | pair | footprint | rank | of | cost | best | "
             "picked | imaging | note |",
             "|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in rows:
        if names:
            design, pair = f"{r.workbook} › {r.design}", r.pair
        else:
            design = (f"W{books.index(r.workbook) + 1}·"
                      f"D{designs.index((r.workbook, r.design)) + 1}")
            pair = f"P{pairs.index((r.workbook, r.design, r.pair)) + 1}"
        lines.append(
            f"| {r.scoring} | {design} | {pair} | {r.footprint} | "
            f"{r.rank if r.rank is not None else '—'} | {r.windows} | "
            f"{r.cost if r.cost is not None else '—'} | {r.best_cost} | "
            f"{'yes' if r.picked else 'no'} | {outcomes.get(r.pair, 'unknown')} | "
            f"{r.note} |")
    return "\n".join(lines)
